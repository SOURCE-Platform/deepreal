"""Bring the Word architecture baseline in line with the Blender model."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/research/DeepReal_System_Architecture_v0.1_historical.docx"
OUTPUT = ROOT / "docs/system-architecture/DeepReal_System_Architecture_Source_of_Truth_v0.2.docx"
IMAGE = ROOT / "docs/system-architecture/current-blender-reference.png"


PARAGRAPH_REPLACEMENTS = {
    "SOURCE OF TRUTH — ENGINEERING BASELINE v0.1":
        "SOURCE OF TRUTH — ENGINEERING BASELINE v0.2",
    "Figure 2 — Mechanical packaging concept. Dimensions shown are architecture-study envelopes, not final industrial-design dimensions.":
        "Figure 2 — Single Blender-authoritative reference concept. The curved lower electronics compartment replaces the colliding Current study and the boxy Compact comparison.",
    "The physical baseline between projector and IR receiver should be as large as practical within the ≈60 mm drum length. The current industrial design places the depth optic and projector group apart along the drum axis. Exact baseline, FOV overlap and algorithm performance are still OPEN and must be determined on an optical bench rather than from CAD alone.":
        "The physical baseline between projector and IR receiver should be as large as practical within the 54 mm drum length. The current industrial design places the depth optic and projector group apart along the drum axis. Exact baseline, FOV overlap and algorithm performance are still OPEN and must be determined on an optical bench rather than from the Blender concept alone.",
    "No slip ring is required because each head is expected to rotate only about 140–160°. Use controlled flex loops through the pivots. High-speed image data should use MIPI-rated FPC/micro-coax, with separate or carefully designed conductors for emitter current, motor power, encoder/control and sensor control. The flexes are a reliability-critical component and must be cycle-tested.":
        "No slip ring is required because each head is expected to rotate only about 140–160°. Use controlled flex loops through the pivots. The Blender model now shows four camera service loops plus two projector power/control routes; these are routing concepts, not production harness drawings. High-speed image data should use MIPI-rated FPC/micro-coax. The flexes remain reliability-critical and must be cycle-tested.",
    "The most important heat sources are the application SoC/PMIC, active IR emitters and regulator losses. Keep the main SoC and power conversion stationary so the body can act as a heat spreader. The rotating projectors need an explicit heat path into the drum carrier/shell and must be duty-cycle limited until measured. Final enclosure material and heat-spreader geometry are OPEN.":
        "The most important heat sources are the application SoC/PMIC, active IR emitters and regulator losses. Keep the main SoC and power conversion stationary. The current Blender concept uses a copper spreader and compliant pad from the PCBA stack to the inside rear housing surface so the enclosure can spread heat. The rotating projectors still need an explicit heat path into the drum carrier/shell and must be duty-cycle limited until measured. Final materials, contact pressure, interface thickness, external temperature and any need for venting remain OPEN.",
}


def _replace_paragraph(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    paragraph.add_run(text)


def _find_table(document, first_header):
    for table in document.tables:
        if table.cell(0, 0).text.strip() == first_header:
            return table
    raise KeyError(first_header)


def _row(table, key):
    for row in table.rows:
        if row.cells[0].text.strip() == key:
            return row
    raise KeyError(key)


def _set_row(row, values):
    for cell, value in zip(row.cells, values):
        cell.text = value


def main():
    document = Document(SOURCE)
    document.tables[0].cell(1, 1).text = "0.2"
    document.tables[0].cell(2, 1).text = "5 September 2026"
    document.tables[0].cell(5, 1).text = (
        "v0.1 architecture baseline plus the reviewed Blender-first "
        "mechanical and packaging update"
    )
    document.tables[1].cell(0, 0).text = (
        "Normative-use rule\nThis document is the architecture and requirements "
        "source of truth. Blender is the current mechanical and visual geometry "
        "authority. When prose and reviewed Blender geometry disagree on "
        "dimensions or placement, update the prose from Blender; architecture "
        "and safety requirements still govern the model. OPEN items must not be "
        "silently converted into assumptions."
    )

    for paragraph in document.paragraphs:
        replacement = PARAGRAPH_REPLACEMENTS.get(paragraph.text.strip())
        if replacement:
            _replace_paragraph(paragraph, replacement)

    enclosure = _find_table(document, "Parameter")
    _set_row(_row(enclosure, "Architecture-study envelope"), (
        "Current Blender width", "120 mm",
        "LOCKED concept width; tolerance-driven CAD remains future work."))
    _set_row(_row(enclosure, "Historical envelope"), (
        "Upper body cross-section", "30 mm H × 24 mm D",
        "Current exterior around the two Ø24 mm drums."))
    _set_row(enclosure.add_row(), (
        "Curved electronics compartment",
        "Extends to ≈29 mm below upper body and ≈24.5 mm rearward",
        "REFERENCE Blender profile; integrated compartment, not a boxy chin."))

    drum = next(t for t in document.tables
                if t.cell(0, 0).text.strip() == "Item"
                and any("Structured-light drum" in r.cells[0].text
                        for r in t.rows))
    _set_row(_row(drum, "Structured-light drum inner diameter"), (
        "Drum outside / assumed inside diameter", "24 mm OD / 21 mm usable ID",
        "Current Blender shell; 1.5 mm architecture-stage wall assumption."))
    _set_row(_row(drum, "Axial length"), (
        "Axial length", "54 mm", "Current Blender shell."))

    optics_values = _find_table(document, "CAD parameter")
    optics_values.cell(0, 0).text = "Blender parameter"

    mechanics = next(t for t in document.tables
                     if t.cell(0, 0).text.strip() == "Item"
                     and any("Keep-out per actuator" in r.cells[0].text
                             for r in t.rows))
    _set_row(_row(mechanics, "Keep-out per actuator"), (
        "Current modeled actuator",
        "Offset cylindrical micro-gearmotor envelope, one per drum; exact part remains OPEN."))

    bom = _find_table(document, "Qty")
    _set_row(next(r for r in bom.rows if r.cells[1].text.strip() == "Motor"), (
        "2", "Motor", "Offset micro geared DC motor",
        "MODELED reference envelope / exact part OPEN", "Stationary behind drums"))
    _set_row(next(r for r in bom.rows
                  if r.cells[1].text.strip() == "Absolute drum encoder"), (
        "2", "Absolute drum encoder + magnet", "AS5600-class magnetic encoder",
        "MODELED reference / exact part OPEN", "Axis"))
    _set_row(next(r for r in bom.rows
                  if r.cells[1].text.strip() == "Drum shells/endcaps/bearings"), (
        "2", "Drum shells/endcaps", "Custom mechanical parts",
        "MODELED concept / material OPEN", "Rotating"))
    _set_row(next(r for r in bom.rows
                  if r.cells[1].text.strip() == "Thermal spreader/shielding"), (
        "1", "EMI shield can + thermal path",
        "Five-piece grounded can, copper spreader and interface pad",
        "MODELED concepts / final designs OPEN", "Stationary PCBA"))
    for values in (
        ("2", "Pinion/ring gear pair", "Custom offset-drive torque transfer",
         "MODELED concept / tooth design OPEN", "Drum ends"),
        ("4 + 2", "Drum bearings and axles", "Miniature bearing/shaft system",
         "MODELED reference envelopes / exact parts OPEN", "Drum axes"),
        ("2", "Mechanical travel stop", "Limits intended drum travel",
         "MODELED concept / final geometry OPEN", "Axis mechanism"),
    ):
        _set_row(bom.add_row(), values)

    gates = _find_table(document, "Gate")
    _set_row(_row(gates, "Gate 8 — Mechanical integration"), (
        "Gate 8 — Mechanical integration",
        "Replace Blender concept envelopes with selected parts and "
        "tolerance-driven geometry inside the curved 120 mm-wide concept.",
        "No collisions through full travel, acceptable weight/hinge behavior, "
        "and serviceable assembly."))

    open_items = _find_table(document, "Open item")
    _set_row(_row(open_items, "Final enclosure OD and mass"), (
        "Final enclosure depth/profile and mass",
        "The curved 120 mm-wide Blender enclosure is a communication and "
        "packaging concept.",
        "Selected parts, tolerance stack, thermal/mass measurements and hinge test."))
    projector = _row(open_items, "Exact projector and exterior aperture pattern")
    projector.cells[1].text = projector.cells[1].text.replace("CAD", "Blender")

    superseded = _find_table(document, "Old assumption")
    _set_row(superseded.rows[1], (
        "Colliding 120 × 30 × 24 mm Current housing",
        "SUPERSEDED and removed. Use the single curved Blender reference concept."))
    _set_row(superseded.add_row(), (
        "Boxy downward Compact chin",
        "SUPERSEDED. Useful electronics volume is integrated into the curved lower compartment."))
    _set_row(superseded.add_row(), (
        "≈150 × 45 × 45 mm planning box is the active design",
        "SUPERSEDED. It remains historical planning context only."))

    heading_five = next(p for p in document.paragraphs
                        if p.text.strip() == "5. Sensor-head architecture"
                        and p.style.name == "Heading 1")
    p = heading_five.insert_paragraph_before("4.7 Electronics compartment EMI shield and heat path")
    p.style = "Heading 2"
    heading_five.insert_paragraph_before(
        "The stationary main PCBA sits in the curved lower compartment. One "
        "local board-mounted conductive shield can encloses the i.MX 95, "
        "CrossLink-NX and memory region. Its lid and four perimeter walls "
        "terminate at PCBA ground. There is no second full-width metal wall. "
        "Camera connectors remain outside the can. A separate copper spreader "
        "and compliant pad show the heat path to the inside rear housing surface."
    )

    heading_four = next(p for p in document.paragraphs
                        if p.text.strip() == "4. Mechanical enclosure and current Blender model"
                        and p.style.name == "Heading 1")
    element = heading_four._p.getnext()
    image_paragraph = next(p for p in document.paragraphs if p._p is element)
    image_paragraph.clear()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(IMAGE), width=Inches(6.35))

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace("v0.1", "v0.2")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
